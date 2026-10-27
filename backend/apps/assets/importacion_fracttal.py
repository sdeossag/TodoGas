"""Importacion de activos desde las exportaciones a Excel de Fracttal.

Fracttal entrega dos archivos (ubicaciones y equipos), una hoja "data" cada uno
con los encabezados en la fila 1. Este modulo no conoce HTTP ni consola:

  leer_ubicaciones / leer_equipos  ->  filas normalizadas
  planificar(ubicaciones, equipos) ->  Plan (solo lee la base de datos)
  aplicar(plan, user)              ->  escribe todo en UNA transaccion

La vista y el comando `importar_fracttal` solo orquestan estas tres piezas.
"""

import io
import re
import unicodedata
import uuid
from dataclasses import dataclass, field
from datetime import date, datetime

from django.db import IntegrityError, transaction

from apps.maintenance.models import MaintenancePlan
from apps.maintenance.services import set_asset_plan

from .models import Asset, AssetNode, Hospital

# --------------------------------------------------------------------------
# Textos
# --------------------------------------------------------------------------


class ImportacionError(Exception):
    """El archivo no se puede importar (no es Excel, faltan columnas...)."""


def _sin_tildes(texto):
    descompuesto = unicodedata.normalize("NFKD", texto)
    return "".join(c for c in descompuesto if not unicodedata.combining(c))


def limpiar(valor):
    """Quita la basura de Fracttal: barras y espacios sobrantes al inicio y al
    final, espacios repetidos, el "---" que usa como vacio."""
    if valor is None:
        return ""
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)
    texto = re.sub(r"\s+", " ", str(valor)).strip()
    texto = texto.strip(" /").strip()
    texto = re.sub(r"\s+", " ", texto)
    return "" if texto == "---" else texto


def normalizar(valor):
    """Clave de comparacion: sin tildes, minusculas, espacios colapsados."""
    texto = _sin_tildes(limpiar(valor)).lower()
    return re.sub(r"\s+", " ", texto).strip()


def _clave_encabezado(valor):
    texto = _sin_tildes("" if valor is None else str(valor)).lower()
    texto = texto.replace("(c)", " ")
    return re.sub(r"\s+", " ", texto).strip()


def partir_ruta(ruta):
    """`// CLINICA NORTE/ IMAGENES/` -> ['CLINICA NORTE', 'IMAGENES']."""
    if not ruta:
        return []
    return [p for p in (limpiar(x) for x in str(ruta).split("/")) if p]


def _id_fracttal(url):
    coincide = re.search(r"/(\d+)\s*/?\s*$", str(url or "").strip())
    return coincide.group(1) if coincide else ""


# --------------------------------------------------------------------------
# Lectura de los Excel
# --------------------------------------------------------------------------

_COLUMNAS = {
    "habilitado": "habilitado",
    "fuera de servicio": "fuera_servicio",
    "descripcion": "descripcion",
    "nombre": "nombre",
    "direccion": "direccion",
    "ciudad": "ciudad",
    "codigo": "codigo",
    "codigo de barras / nfc": "barcode",
    "prioridad": "prioridad",
    "tipo": "tipo",
    "clasificacion 1": "clasificacion_1",
    "clasificacion 2": "clasificacion_2",
    "proveedor": "proveedor",
    "fecha de compra": "fecha_compra",
    "plan de tareas": "plan",
    "notas": "notas",
    "fabricante": "fabricante",
    "modelo": "modelo",
    "numero de serial": "serial",
    "ubicacion de equipo": "ubicacion_equipo",
    "url codigo qr": "url_qr",
}
_CRUDOS = {"fecha_compra"}  # no se limpian como texto


@dataclass
class Fila:
    archivo: str  # "ubicaciones" | "equipos"
    fila: int  # numero de fila de Excel (encabezado = 1)
    datos: dict

    @property
    def nombre(self):
        return self.datos.get("nombre", "")


def _abrir(archivo, tipo):
    from openpyxl import load_workbook

    try:
        if hasattr(archivo, "read"):
            if hasattr(archivo, "seek"):
                archivo.seek(0)
            contenido = io.BytesIO(archivo.read())
        else:
            contenido = archivo
        return load_workbook(contenido, read_only=True, data_only=True)
    except Exception as exc:
        raise ImportacionError(
            f"El archivo de {tipo} no es un Excel (.xlsx) válido."
        ) from exc


def _leer(archivo, tipo):
    wb = _abrir(archivo, tipo)
    try:
        hoja = wb["data"] if "data" in wb.sheetnames else wb.worksheets[0]
        filas = hoja.iter_rows(values_only=True)
        try:
            encabezado = next(filas)
        except StopIteration:
            raise ImportacionError(f"El archivo de {tipo} está vacío.")

        columnas = {}
        for pos, titulo in enumerate(encabezado):
            clave = _clave_encabezado(titulo)
            if clave.startswith("ubicado en"):
                destino = "ruta"
            else:
                destino = _COLUMNAS.get(clave)
            if destino and destino not in columnas:
                columnas[destino] = pos

        faltan = [
            etiqueta
            for destino, etiqueta in (("nombre", "Nombre"), ("ruta", "Ubicado en ó es Parte de"))
            if destino not in columnas
        ]
        if faltan:
            raise ImportacionError(
                f"Al archivo de {tipo} le faltan columnas obligatorias: {', '.join(faltan)}."
            )

        resultado = []
        for numero, celdas in enumerate(filas, start=2):
            if not celdas or all(c is None or str(c).strip() == "" for c in celdas):
                continue
            datos = {}
            for destino, pos in columnas.items():
                valor = celdas[pos] if pos < len(celdas) else None
                datos[destino] = valor if destino in _CRUDOS else limpiar(valor)
            datos["ruta"] = (
                "" if celdas[columnas["ruta"]] is None else str(celdas[columnas["ruta"]])
            )
            datos["fracttal_id"] = _id_fracttal(datos.get("url_qr", ""))
            resultado.append(Fila(tipo, numero, datos))
        return resultado
    finally:
        wb.close()


def leer_ubicaciones(archivo):
    return _leer(archivo, "ubicaciones")


def leer_equipos(archivo):
    return _leer(archivo, "equipos")


# --------------------------------------------------------------------------
# Plan
# --------------------------------------------------------------------------

_SI = {"si", "yes", "true", "1", "s"}
_NO = {"no", "false", "0", "n"}
_PRIORIDAD = {
    "alta": Asset.Priority.HIGH,
    "media": Asset.Priority.MEDIUM,
    "baja": Asset.Priority.LOW,
}


@dataclass
class HospitalPlan:
    nombre: str
    existente: object = None  # Hospital | None
    codigo: str = ""  # solo si es nuevo


@dataclass
class NodoPlan:
    hospital_key: str
    nombres: tuple
    claves: tuple
    existente: object = None  # AssetNode | None
    explicito: bool = False  # venia en el archivo de ubicaciones
    fracttal_id: str = ""
    codigo: str = ""
    direccion: str = ""
    activo: bool = True
    fila_origen: dict = None  # fila que lo pidio sin que viniera en el archivo

    @property
    def accion(self):
        return "update" if self.existente is not None else "create"


@dataclass
class ActivoPlan:
    hospital_key: str
    nodo_key: tuple = None  # (hospital_key, claves) o None si cuelga del hospital
    existente: object = None  # Asset | None
    campos: dict = field(default_factory=dict)
    fracttal_id: str = ""
    plan: object = None  # MaintenancePlan | None (ya existente)
    plan_nuevo: str = ""  # clave normalizada de un protocolo por crear


@dataclass
class Plan:
    hospitales: dict = field(default_factory=dict)
    nodos: dict = field(default_factory=dict)
    activos: list = field(default_factory=list)
    filas: list = field(default_factory=list)
    planes_crear: dict = field(default_factory=dict)  # clave normalizada -> nombre
    planes_existentes: list = field(default_factory=list)  # nombres

    @property
    def errores(self):
        return sum(len(f["errors"]) for f in self.filas)

    @property
    def avisos(self):
        return sum(len(f["warnings"]) for f in self.filas)

    def resumen(self, aplicado=False):
        crear = [h.nombre for h in self.hospitales.values() if h.existente is None]
        existentes = [h.nombre for h in self.hospitales.values() if h.existente is not None]
        nodos_crear = sum(1 for n in self.nodos.values() if n.existente is None)
        nodos_act = sum(1 for n in self.nodos.values() if n.explicito and n.existente is not None)
        return {
            "hospitals": {"create": crear, "existing": existentes},
            "locations": {"create": nodos_crear, "update": nodos_act},
            "assets": {
                "create": sum(1 for a in self.activos if a.existente is None),
                "update": sum(1 for a in self.activos if a.existente is not None),
            },
            "plans": {
                "create": list(self.planes_crear.values()),
                "existing": list(self.planes_existentes),
            },
            "rows": self.filas,
            "errors_count": self.errores,
            "warnings_count": self.avisos,
            "applied": aplicado,
        }


def _fecha(valor):
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    texto = limpiar(valor)
    if not texto:
        return None
    for formato in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%Y/%m/%d"):
        try:
            return datetime.strptime(texto[:10], formato).date()
        except ValueError:
            continue
    raise ValueError(texto)


def _slug(texto, largo):
    base = re.sub(r"[^A-Z0-9]+", "-", _sin_tildes(texto).upper()).strip("-")
    return base[:largo].strip("-") or "SIN-NOMBRE"


def _unico(base, usados, largo):
    """`base`, o `base-2`, `base-3`... que no este en `usados` (minusculas)."""
    base = base[:largo]
    candidato, n = base, 1
    while candidato.lower() in usados:
        n += 1
        sufijo = f"-{n}"
        candidato = base[: largo - len(sufijo)] + sufijo
    usados.add(candidato.lower())
    return candidato


class _Planificador:
    def __init__(self):
        self.plan = Plan()
        self.hospitales_db = {}
        for h in Hospital.objects.order_by("created_at"):
            self.hospitales_db.setdefault(normalizar(h.name), h)
        self.codigos_hospital = {c.lower() for c in Hospital.objects.values_list("code", flat=True)}
        self.codigos_activo = {c.lower() for c in Asset.objects.values_list("code", flat=True)}
        self.nodos_db = {}
        self.nodos_por_fid = {}
        for n in AssetNode.objects.all():
            self.nodos_db[(n.hospital_id, n.parent_id, normalizar(n.name))] = n
            if n.fracttal_id:
                self.nodos_por_fid[n.fracttal_id] = n
        self.activos_por_fid = {
            a.fracttal_id: a
            for a in Asset.objects.exclude(fracttal_id="").select_related("plan")
        }
        self.planes_db = {}
        for p in MaintenancePlan.objects.order_by("created_at"):
            self.planes_db.setdefault(normalizar(p.name), p)
        self.fids_nodos, self.fids_activos = set(), set()
        self.planes_existentes = {}

    # --- filas de salida ---------------------------------------------------
    def _fila(self, f, ubicacion=""):
        salida = {
            "file": f.archivo,
            "row": f.fila,
            "name": f.nombre,
            "location": ubicacion,
            "action": "create",
            "errors": [],
            "warnings": [],
        }
        self.plan.filas.append(salida)
        return salida

    @staticmethod
    def _error(salida, mensaje):
        salida["errors"].append(mensaje)
        salida["action"] = "error"

    # --- hospitales y nodos -------------------------------------------------
    def _hospital(self, nombre):
        clave = normalizar(nombre)
        hp = self.plan.hospitales.get(clave)
        if hp is None:
            existente = self.hospitales_db.get(clave)
            hp = HospitalPlan(nombre=existente.name if existente else nombre, existente=existente)
            if existente is None:
                hp.codigo = _unico(_slug(nombre, 40), self.codigos_hospital, 50)
            self.plan.hospitales[clave] = hp
        return hp

    def _buscar_nodo(self, hospital, padre, nombre):
        if hospital.existente is None:
            return None
        if padre is not None and padre.existente is None:
            return None
        return self.nodos_db.get(
            (hospital.existente.id, padre.existente.id if padre else None, normalizar(nombre))
        )

    def _asegurar(self, hk, hospital, nombres, fila):
        """Garantiza que existan (en el plan) todas las ubicaciones de la ruta.
        Devuelve la ultima, o None si la ruta es solo el hospital."""
        padre = None
        for i, nombre in enumerate(nombres):
            claves = tuple(normalizar(n) for n in nombres[: i + 1])
            clave = (hk, claves)
            nodo = self.plan.nodos.get(clave)
            if nodo is None:
                nodo = NodoPlan(
                    hospital_key=hk,
                    nombres=tuple(nombres[: i + 1]),
                    claves=claves,
                    existente=self._buscar_nodo(hospital, padre, nombre),
                    fila_origen=fila,
                )
                self.plan.nodos[clave] = nodo
            padre = nodo
        return padre

    # --- ubicaciones --------------------------------------------------------
    def ubicaciones(self, filas):
        for f in filas:
            d = f.datos
            partes = partir_ruta(d.get("ruta", ""))
            salida = self._fila(f, " / ".join(partes))
            if not f.nombre:
                self._error(salida, "La fila no tiene nombre.")
            if not partes:
                self._error(salida, "La ruta «Ubicado en» no indica el hospital.")
            if salida["errors"]:
                continue

            fid = d.get("fracttal_id", "")
            if fid:
                if fid in self.fids_nodos:
                    self._error(salida, f"El id de Fracttal {fid} está repetido en el archivo.")
                    continue
                self.fids_nodos.add(fid)

            hk = normalizar(partes[0])
            por_fid = self.nodos_por_fid.get(fid) if fid else None
            if por_fid is not None:
                hospital_db = self.hospitales_db.get(hk)
                if hospital_db is None or por_fid.hospital_id != hospital_db.id:
                    self._error(
                        salida,
                        f"El id de Fracttal {fid} ya pertenece a una ubicación de otro hospital.",
                    )
                    continue

            hospital = self._hospital(partes[0])
            padre = self._asegurar(hk, hospital, partes[1:], salida)
            nombres = partes[1:] + [f.nombre]
            clave = (hk, tuple(normalizar(n) for n in nombres))
            nodo = self.plan.nodos.get(clave)
            if nodo is not None and nodo.explicito:
                self._error(salida, "Esta ubicación está repetida en el archivo.")
                continue
            if nodo is None:
                nodo = NodoPlan(
                    hospital_key=hk,
                    nombres=tuple(nombres),
                    claves=clave[1],
                    existente=self._buscar_nodo(hospital, padre, f.nombre),
                )
                self.plan.nodos[clave] = nodo
            nodo.fila_origen = None  # ya no es intermedia
            nodo.explicito = True

            if por_fid is not None:
                if nodo.existente is not None and nodo.existente.id != por_fid.id:
                    salida["warnings"].append(
                        "Ya existía otra ubicación con ese nombre en esa ruta; "
                        "se actualiza la que tiene el mismo id de Fracttal."
                    )
                nodo.existente = por_fid
                padre_id = padre.existente.id if padre and padre.existente else None
                if (
                    padre is not None and padre.existente is None
                ) or por_fid.parent_id != padre_id or normalizar(por_fid.name) != normalizar(f.nombre):
                    salida["warnings"].append(
                        "La ubicación ya existe con otro nombre o en otro lugar; no se mueve."
                    )
            elif nodo.existente is not None and fid and nodo.existente.fracttal_id not in ("", fid):
                salida["warnings"].append(
                    f"La ubicación ya tenía el id de Fracttal {nodo.existente.fracttal_id}; "
                    f"se reemplaza por {fid}."
                )

            nodo.fracttal_id = fid
            nodo.codigo = d.get("codigo", "")[:50]
            direccion = ", ".join(x for x in (d.get("direccion", ""), d.get("ciudad", "")) if x)
            nodo.direccion = direccion[:255]
            nodo.activo = normalizar(d.get("habilitado", "")) not in _NO
            salida["action"] = nodo.accion

    # --- equipos ------------------------------------------------------------
    def equipos(self, filas):
        for f in filas:
            d = f.datos
            partes = partir_ruta(d.get("ruta", ""))
            salida = self._fila(f, " / ".join(partes))
            if not f.nombre:
                self._error(salida, "La fila no tiene nombre.")
            if not partes:
                self._error(salida, "La ruta «Ubicado en» no indica el hospital.")
            if salida["errors"]:
                continue

            fid = d.get("fracttal_id", "")
            if fid:
                if fid in self.fids_activos:
                    self._error(salida, f"El id de Fracttal {fid} está repetido en el archivo.")
                    continue
                self.fids_activos.add(fid)
            else:
                salida["warnings"].append(
                    "La fila no trae id de Fracttal (Url Código QR): al reimportar se duplicará."
                )

            hk = normalizar(partes[0])
            hospital = self._hospital(partes[0])
            nodo = self._asegurar(hk, hospital, partes[1:], salida)
            existente = self.activos_por_fid.get(fid) if fid else None

            campos = {"name": f.nombre[:255]}
            if len(f.nombre) > 255:
                salida["warnings"].append("El nombre se recortó a 255 caracteres.")
            limites = (
                ("manufacturer", "fabricante", 200),
                ("model", "modelo", 200),
                ("serial_number", "serial", 200),
                ("equipment_location", "ubicacion_equipo", 255),
                ("barcode", "barcode", 100),
                ("asset_type", "tipo", 100),
                ("classification_1", "clasificacion_1", 100),
                ("classification_2", "clasificacion_2", 100),
                ("supplier", "proveedor", 200),
                ("notes", "notas", None),
            )
            for campo, origen, largo in limites:
                valor = d.get(origen, "")
                if largo and len(valor) > largo:
                    salida["warnings"].append(f"«{origen}» se recortó a {largo} caracteres.")
                    valor = valor[:largo]
                campos[campo] = valor

            prioridad = normalizar(d.get("prioridad", ""))
            if prioridad and prioridad not in _PRIORIDAD:
                salida["warnings"].append(f"Prioridad «{d['prioridad']}» no reconocida; queda en Media.")
            campos["priority"] = _PRIORIDAD.get(prioridad, Asset.Priority.MEDIUM)

            if normalizar(d.get("habilitado", "")) in _NO:
                campos["status"] = Asset.Status.DECOMMISSIONED
            elif normalizar(d.get("fuera_servicio", "")) in _SI:
                campos["status"] = Asset.Status.OUT_OF_SERVICE
            else:
                campos["status"] = Asset.Status.ACTIVE

            try:
                campos["purchase_date"] = _fecha(d.get("fecha_compra"))
            except ValueError:
                campos["purchase_date"] = None
                salida["warnings"].append("Fecha de compra no válida; se ignora.")

            if existente is not None:
                campos["code"] = existente.code
            else:
                deseado = d.get("codigo", "")
                if deseado and len(deseado) <= 50 and deseado.lower() not in self.codigos_activo:
                    self.codigos_activo.add(deseado.lower())
                    campos["code"] = deseado
                else:
                    base = "FRC-" + (fid or _slug(f.nombre, 30))
                    campos["code"] = _unico(base, self.codigos_activo, 50)
                    if not deseado:
                        salida["warnings"].append(
                            f"La fila no trae código; se generó «{campos['code']}»."
                        )
                    else:
                        motivo = "es demasiado largo" if len(deseado) > 50 else "ya lo usa otro activo"
                        salida["warnings"].append(
                            f"El código «{deseado}» {motivo}; se generó «{campos['code']}»."
                        )

            nombre_plan = d.get("plan", "")[:255]
            plan_obj = None
            plan_nuevo = ""
            if nombre_plan:
                clave_plan = normalizar(nombre_plan)
                plan_obj = self.planes_db.get(clave_plan)
                if plan_obj is not None:
                    self.planes_existentes.setdefault(clave_plan, plan_obj.name)
                    destino = plan_obj.name
                else:
                    plan_nuevo = clave_plan
                    self.plan.planes_crear.setdefault(clave_plan, nombre_plan)
                    destino = self.plan.planes_crear[clave_plan]
                    salida["warnings"].append(
                        f"Se creará el protocolo «{destino}» vacío; "
                        "hay que completarle tareas y checklist."
                    )
                actual = existente.plan if existente is not None else None
                if actual is not None and (plan_obj is None or actual.id != plan_obj.id):
                    salida["warnings"].append(
                        f"Cambia de protocolo: «{actual.name}» → «{destino}»."
                    )

            self.plan.activos.append(
                ActivoPlan(
                    hospital_key=hk,
                    nodo_key=(hk, nodo.claves) if nodo else None,
                    existente=existente,
                    campos=campos,
                    fracttal_id=fid,
                    plan=plan_obj,
                    plan_nuevo=plan_nuevo,
                )
            )
            salida["action"] = "update" if existente else "create"

    def cerrar(self):
        for nodo in self.plan.nodos.values():
            if not nodo.explicito and nodo.existente is None and nodo.fila_origen is not None:
                nodo.fila_origen["warnings"].append(
                    f"La ubicación intermedia «{nodo.nombres[-1]}» no venía en el archivo "
                    "de ubicaciones; se crea igual."
                )
        self.plan.planes_existentes = list(self.planes_existentes.values())
        return self.plan


def planificar(ubicaciones, equipos):
    """Arma el plan de importacion. Solo lee la base de datos."""
    p = _Planificador()
    p.ubicaciones(ubicaciones or [])
    p.equipos(equipos or [])
    return p.cerrar()


# --------------------------------------------------------------------------
# Aplicar
# --------------------------------------------------------------------------


DESCRIPCION_PROTOCOLO_NUEVO = (
    "Creado desde la importación de Fracttal. Falta agregar tareas y checklist."
)


def aplicar(plan, user):
    """Escribe el plan en una sola transaccion. Devuelve los conteos aplicados."""
    if plan.errores:
        raise ImportacionError("El plan tiene errores; no se puede aplicar.")

    conteos = {
        "hospitals_created": 0,
        "locations_created": 0,
        "locations_updated": 0,
        "assets_created": 0,
        "assets_updated": 0,
        "plans_created": 0,
        "plans_assigned": 0,
    }
    try:
        with transaction.atomic():
            hospitales = {}
            for clave, hp in plan.hospitales.items():
                if hp.existente is not None:
                    hospitales[clave] = hp.existente
                else:
                    hospitales[clave] = Hospital.objects.create(name=hp.nombre, code=hp.codigo)
                    conteos["hospitals_created"] += 1

            planes_nuevos = {}
            for clave, nombre in plan.planes_crear.items():
                planes_nuevos[clave] = MaintenancePlan.objects.create(
                    name=nombre, description=DESCRIPCION_PROTOCOLO_NUEVO
                )
                conteos["plans_created"] += 1

            nodos = {}
            for clave, n in sorted(plan.nodos.items(), key=lambda kv: len(kv[1].claves)):
                hospital = hospitales[n.hospital_key]
                if n.existente is not None:
                    obj = AssetNode.objects.get(pk=n.existente.pk)
                    if n.explicito:
                        if n.fracttal_id:
                            obj.fracttal_id = n.fracttal_id
                        if n.codigo:
                            obj.code = n.codigo
                        if n.direccion:
                            obj.address = n.direccion
                        obj.is_active = n.activo
                        obj.save()
                        conteos["locations_updated"] += 1
                else:
                    padre = nodos.get((n.hospital_key, n.claves[:-1])) if len(n.claves) > 1 else None
                    obj = AssetNode(
                        hospital=hospital,
                        parent=padre,
                        name=n.nombres[-1],
                        code=n.codigo,
                        address=n.direccion,
                        is_active=n.activo,
                        fracttal_id=n.fracttal_id,
                    )
                    obj.save()
                    conteos["locations_created"] += 1
                nodos[clave] = obj

            for a in plan.activos:
                hospital = hospitales[a.hospital_key]
                nodo = nodos.get(a.nodo_key) if a.nodo_key else None
                if a.existente is not None:
                    activo = Asset.objects.get(pk=a.existente.pk)
                    for campo, valor in a.campos.items():
                        # Un dato vacio de Fracttal no borra lo que ya se
                        # completo aqui; el codigo no cambia al reimportar.
                        if campo == "code" or valor in ("", None):
                            continue
                        setattr(activo, campo, valor)
                    activo.hospital = hospital
                    activo.node = nodo
                    activo.save()
                    conteos["assets_updated"] += 1
                else:
                    activo = Asset(
                        hospital=hospital,
                        node=nodo,
                        fracttal_id=a.fracttal_id,
                        **a.campos,
                    )
                    activo.save()
                    conteos["assets_created"] += 1
                destino = a.plan if a.plan is not None else planes_nuevos.get(a.plan_nuevo)
                if destino is not None and activo.plan_id != destino.id:
                    set_asset_plan(activo, destino, created_by=user)
                    conteos["plans_assigned"] += 1
    except IntegrityError as exc:
        raise ImportacionError(f"La base de datos rechazó la importación: {exc}") from exc
    return conteos


# --------------------------------------------------------------------------
# Auditoria
# --------------------------------------------------------------------------


def auditar(user, archivos, resumen, conteos, ip=None, user_agent=""):
    """Deja constancia de la importacion (quien, que archivos, cuantos)."""
    from apps.audit.models import AuditLog

    AuditLog.objects.create(
        user=user,
        action=AuditLog.Action.CREATE,
        entity_type="FracttalImport",
        entity_id=uuid.UUID(int=0),
        changes={
            "archivos": archivos,
            "hospitales_nuevos": len(resumen["hospitals"]["create"]),
            "ubicaciones": resumen["locations"],
            "equipos": resumen["assets"],
            "protocolos_nuevos": len(resumen["plans"]["create"]),
            "avisos": resumen["warnings_count"],
            "aplicado": conteos,
        },
        ip_address=ip or None,
        user_agent=(user_agent or "")[:500],
    )
